"""Persistent client/session organization around immutable metrological results."""
from __future__ import annotations

import csv
import io
import json
import zipfile

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from ..domain.processing.export import (_build_client_filename, _sanitize_filename, _build_client_output_frame,
                                       build_client_output_workbook_bytes, build_dataset_workbook_bytes)
from .reports import build_session_dossier

from .models import Decision
from .repository import encode
from .importer import measurement_identity


class ResultsSessions:
    def set_workbook_inclusion(self, session_id, run_id, included, command):
        """Change consultation membership while preserving imports and evaluations."""
        with self.repo.connect(write=True) as db:
            session = self.repo.get(db, "results_sessions", session_id)
            active = list(session["run_ids"])
            detached = list(session.get("detached_run_ids", []))
            if run_id not in active + detached:
                raise ValueError("Choose a workbook imported into this session")
            if included and run_id in detached:
                detached.remove(run_id)
                active.append(run_id)
            elif not included and run_id in active:
                active.remove(run_id)
                detached.append(run_id)
            session.update(run_ids=active, detached_run_ids=detached)
            self.repo.update(db, "results_sessions", session_id, session)
            self.repo.audit(db, "session_workbook_inclusion", session_id, command.actor, command.reason,
                            after={"run_id": run_id, "included": included})
            return session

    def save_outlier_screening(self, session_id, command):
        from .qc_screening import stored_qc_screening
        with self.repo.connect(write=True) as db:
            session = self.repo.get(db, "results_sessions", session_id)
            settings = command.model_dump(exclude={"actor", "reason"})
            before = session.get("outlier_screening", {"method": "sigma", "threshold": 3.0})
            session["outlier_screening"] = settings
            self.repo.update(db, "results_sessions", session_id, session)
            result = stored_qc_screening(self.repo, db, session)
            if before != settings:
                self.repo.audit(db, "qc_screening_settings_saved", session_id, command.actor, command.reason, before, settings)
            return result

    def results_session_analysis(self, session_id, outlier_method=None, threshold=None, range_exclusions=None, include_all_data=False):
        from .qc_screening import stored_qc_screening
        from .session_analysis import session_analysis
        detail = self.results_session_detail(session_id)
        # Linearity uses the first acquired cycle, as in the original IRMS charts.
        # Evaluation and uncertainty records retain their recorded intensity basis.
        from .analysis_evidence import analysis_evidence
        with self.repo.connect() as db:
            for run in detail["runs"]:
                initial = {}
                for record in self.repo.list(db, "measurements", run_id=run["id"]):
                    cycles = analysis_evidence(record, run["source"])["cycles"]
                    if cycles:
                        initial[record["id"]] = cycles[0]
                for row in (run.get("evaluation") or {}).get("results", []):
                    cycle = initial.get(row["id"], {})
                    row["initial_i44_v"] = cycle.get("i44_v", row.get("i44_v"))
                    row["initial_reference_i44_v"] = cycle.get("reference_i44_v", row.get("reference_i44_v"))
        # Explicit query parameters remain a non-persistent preview for existing clients.
        if outlier_method is not None or threshold is not None:
            settings = detail.get("outlier_screening", {"method": "sigma", "threshold": 3.0})
            return session_analysis(detail, outlier_method=outlier_method or settings["method"],
                                    threshold=threshold if threshold is not None else settings["threshold"], range_exclusions=range_exclusions, include_all_data=include_all_data)
        with self.repo.connect(write=True) as db:
            session = self.repo.get(db, "results_sessions", session_id)
            outliers = stored_qc_screening(self.repo, db, session, detail["runs"])
        return session_analysis(detail, outliers=outliers, range_exclusions=range_exclusions, include_all_data=include_all_data)

    def save_chart_settings(self, session_id, command):
        with self.repo.connect(write=True) as db:
            session = self.repo.get(db, "results_sessions", session_id)
            before = session.get("chart_settings", {})
            if command.diagnostic_material_id and command.diagnostic_material_id not in self.material_map(db, self.repo.get(db, "methods", session["method_id"])):
                raise ValueError("Choose a material belonging to the session method")
            session["chart_settings"] = command.model_dump(exclude={"actor", "reason"})
            self.repo.update(db, "results_sessions", session_id, session)
            self.repo.audit(db, "session_chart_settings_saved", session_id, command.actor, command.reason, before, session["chart_settings"])
            return session

    def save_residual_override(self, session_id, command):
        with self.repo.connect(write=True) as db:
            session = self.repo.get(db, "results_sessions", session_id)
            if command.material_id != "__all__" and command.material_id not in self.material_map(db, self.repo.get(db, "methods", session["method_id"])):
                raise ValueError("Choose a material belonging to the session method")
            key = f"{command.material_id}:{command.effect}:{command.isotope}"
            overrides = session.setdefault("residual_overrides", {})
            before = overrides.get(key)
            if command.settings is None:
                overrides.pop(key, None)
            else:
                overrides[key] = command.settings.model_dump()
            self.repo.update(db, "results_sessions", session_id, session)
            self.repo.audit(db, "residual_preview_override_saved", session_id, command.actor, command.reason,
                            {"key": key, "settings": before}, {"key": key, "settings": overrides.get(key), "scope": "review_preview"})
            return session

    def save_results_session(self, command, session_id=None):
        if command.calibration_verification=="simulation_assumption" and not self.repo.demo:
            raise ValueError("Simulated qualification assumptions require a separate demonstration workspace")
        with self.repo.connect(write=True) as db:
            method=self.repo.get(db,"methods",command.method_id)
            previous=self.repo.get(db,"results_sessions",session_id) if session_id else None
            if previous and previous["run_ids"]:
                # Session identity and processing assumptions form part of every result's provenance.
                raise ValueError("An imported session is fixed. Create a new session for a different client, method or processing basis.")
            if command.context=="routine" and method["status"]!="active":
                raise ValueError("A new routine session requires the current validated method")
            qid=command.qualification_id
            if command.context=="qualification":
                if not qid:
                    raise ValueError("Choose an open qualification session")
                q=self.check_open_qualification(db,qid)
                if q["method_id"]!=command.method_id:
                    raise ValueError("Qualification and method must match")
                if command.input_basis!="instrument_delta" or command.preapplied_corrections:
                    raise ValueError("New qualification uses observations before secondary correction and normalization")
            else:
                qid=method.get("approval",{}).get("qualification_id")
                if not qid:
                    raise ValueError("The method needs its approved qualification")
            data=command.model_dump(exclude={"method_id","actor","reason"}) | {
                "qualification_id":qid,"run_ids":[],"groups":{},"bridges":{},"simulation":self.repo.demo,
                "method_name":command.method_name or method["config"]["name"],
                "intended_use":command.intended_use or method["config"]["intended_use"],
            }
            if previous:
                if previous["method_id"]!=command.method_id:
                    db.execute("UPDATE results_sessions SET method_id=? WHERE id=?",(command.method_id,session_id))
                self.repo.update(db,"results_sessions",session_id,data)
                record=self.repo.get(db,"results_sessions",session_id)
            else:
                record=self.repo.insert(db,"results_sessions",data,method_id=command.method_id)
            self.repo.audit(db,"results_session_saved",record["id"],command.actor,command.reason,previous,record)
            return record

    def prepare_session_import(self, db, command):
        if not command.results_session_id:
            return None
        session=self.repo.get(db,"results_sessions",command.results_session_id)
        if command.method_id and command.method_id!=session["method_id"]:
            raise ValueError("The workbook must use the method pinned to this results session")
        command.method_id=session["method_id"]
        command.context=session["context"]
        command.qualification_id=session["qualification_id"] if session["context"]=="qualification" else None
        command.input_basis=session["input_basis"]
        command.preapplied_corrections=session["preapplied_corrections"]
        command.processing_evidence=session["processing_evidence"]
        command.external_method_id=session["method_id"] if command.input_basis=="already_vpdb" or command.preapplied_corrections else None
        return session

    def attach_session_run(self, db, session, run, command):
        if run["id"] in session["run_ids"]:
            return
        owner=next((s for s in self.repo.list(db,"results_sessions") if run["id"] in [*s["run_ids"], *s.get("detached_run_ids", [])]),None)
        if owner and owner["id"] != session["id"]:
            raise ValueError(f"Workbook already belongs to results session '{owner['name']}' for client '{owner['client']}'")
        if run["method_id"]!=session["method_id"] or run["context"]!=session["context"]:
            raise ValueError("Existing import has a different method or workflow")
        session["run_ids"].append(run["id"])
        session["detached_run_ids"] = [rid for rid in session.get("detached_run_ids", []) if rid != run["id"]]
        for row in self.repo.list(db,"measurements",run_id=run["id"]):
            session["groups"][row["id"]]=command.sample_group.strip() or "Main batch"
        self.repo.update(db,"results_sessions",session["id"],session)
        self.repo.audit(db,"workbook_attached_to_session",session["id"],command.actor,command.reason,after={"run_id":run["id"],"sample_group":command.sample_group})

    def session_catalog(self, db):
        runs={r["id"]:r for r in self.repo.list(db,"runs")}
        result=[]
        for session in self.repo.list(db,"results_sessions"):
            linked=[runs[r] for r in session["run_ids"]]
            status="empty" if not linked else ("blocked" if any(r["status"]=="blocked" for r in linked) else ("released" if all(r["status"]=="released" for r in linked) else "awaiting_review"))
            if session["context"]=="qualification" and session["qualification_id"]:
                status=self.repo.get(db,"qualifications",session["qualification_id"])["status"]
            result.append({**session,"status":status,"analysis_count":sum(r["analysis_count"] for r in linked),
                           "acquired_date":min((r["acquired_date"] for r in linked if r["acquired_date"]),default=None),
                           "sample_groups":sorted(set(session["groups"].values()) or {"Main batch"})})
        return result

    def results_session_detail(self, session_id):
        with self.repo.connect() as db:
            session=self.repo.get(db,"results_sessions",session_id)
            method=self.repo.get(db,"methods",session["method_id"]) if session["method_id"] else None
            q=self.qualification_detail(db,session["qualification_id"]) if session["qualification_id"] else None
            reference_run=self.repo.get(db,"evaluations",q["approval"]["evaluation_id"])["run_id"] if q and q.get("approval") else None
            history=self.qc_history(db,session["method_id"]) if session["method_id"] else []
            if session.get("calibration_verification")=="simulation_assumption":
                history=[h for h in history if h.get("data_origin")=="observed" and h.get("results_session_id")==session_id]
            sources=[a for a in self.repo.list(db,"session_sources") if a["session_id"]==session_id]
            exports=[e for e in self.repo.list(db,"session_exports") if e["session_id"]==session_id]
        return {**session,"method":method,"qualification":q,"qualification_run_id":reference_run,
                "runs":[self.run_detail(r) for r in session["run_ids"]],
                "detached_runs":[self.run_detail(r) for r in session.get("detached_run_ids", [])],
                "history":history,"exports":exports,"sources":sources}

    def archive_session_source(self, session_id, filename, content, command, *, relative_path, disposition, run_id=None):
        """Archive only original acquisition files, including superseded raw exports."""
        digest=self.repo.blob(content)
        with self.repo.connect(write=True) as db:
            self.repo.get(db,"results_sessions",session_id)
            previous=next((s for s in self.repo.list(db,"session_sources") if s["session_id"]==session_id and s["relative_path"]==relative_path and s["sha256"]==digest),None)
            if previous: return previous
            record=self.repo.insert(db,"session_sources",{"session_id":session_id,"filename":filename,"relative_path":relative_path,
                                    "sha256":digest,"size":len(content),"disposition":disposition,"run_id":run_id})
            self.repo.audit(db,"raw_session_source_archived",session_id,command.actor,command.reason,after=record)
            return record

    def save_session_groups(self, session_id, command):
        with self.repo.connect(write=True) as db:
            session=self.repo.get(db,"results_sessions",session_id)
            measurements={r["id"]:run_id for run_id in session["run_ids"] for r in self.repo.list(db,"measurements",run_id=run_id)}
            if not set(command.groups).issubset(measurements):
                raise ValueError("Only measurements belonging to this session can be grouped")
            if any(not name.strip() or len(name)>160 for name in command.groups.values()):
                raise ValueError("Give each sample group a name of 1–160 characters")
            if any(self.repo.get(db,"runs",measurements[mid])["status"]=="released" for mid in command.groups):
                raise ValueError("Sample groups in released results are frozen")
            before=dict(session["groups"])
            session["groups"].update(command.groups)
            self.repo.update(db,"results_sessions",session_id,session)
            self.repo.audit(db,"session_sample_groups_saved",session_id,command.actor,command.reason,before,session["groups"])
            return session

    def export_results_session(self, session_id, command):
        detail=self.results_session_detail(session_id)
        analysis=self.results_session_analysis(session_id)
        flagged={flag["measurement_id"] for flag in analysis["outliers"]["flags"]}
        rows, whole_rows = [], []
        for run in detail["runs"]:
            evaluation=run["evaluation"]
            if not evaluation:
                continue
            for row in evaluation["results"]:
                group=detail["groups"].get(row["id"],"Main batch")
                if row["role"]=="unknown" and command.group is not None and command.group!=group:
                    continue
                item={"client":detail["client"],"project":detail["project"],"session":detail["name"],"session_id":session_id,
                      "sample_group":group,"sample":row["label"],"analysis":row["source_index"],"run_id":run["id"],
                      "acquired_at":row.get("acquired_at"),"worksheet_row":row.get("sheet_row"),"mass_ug":row.get("mass_ug"),
                      "sample_i44_v":row.get("i44_v"),"pressure_adjustment_difference_v":row.get("pressure_mismatch_v"),
                      "evaluation_id":evaluation["id"],"method_id":run["method_id"],"method_version":detail["method"]["version"],
                      "qualification_id":detail["qualification_id"],"raw_sha256":run["source"]["sha256"],"input_basis":run.get("input_basis"),
                      "decision":"released" if run["status"]=="released" else ("blocked" if row["issues"] or evaluation["blockers"] or run["release_blockers"] else "review"),
                      "issues":"; ".join(dict.fromkeys(row["issues"]+[issue for issue in evaluation["blockers"]+([] if run["status"]=="released" else run["release_blockers"]) if not issue.startswith("Analysis ") ])),"simulation":self.repo.demo,
                      "accepted_exceptions":"; ".join(row.get("accepted_issues", [])),
                      "data_origin":"synthetic" if run.get("synthetic") else "observed","source_kind":run.get("source_kind","qtegra_raw"),
                      "calibration_verification":run.get("calibration_verification","documented"),"sample_identifier":row.get("comment","")}
                item.update(measurement_identity(row, run.get("source_kind", "qtegra_raw")))
                item.update(measurement_id=row["id"],role=row["role"], excluded=bool(row["excluded"] or row["id"] in flagged),
                            outlier_isotopes=", ".join(f["isotope"] for f in analysis["outliers"]["flags"] if f["measurement_id"]==row["id"]))
                if item["excluded"]: item["decision"]="excluded"
                item.update(d13c_internal_sd=row.get("d13c_sd"), d18o_internal_sd=row.get("d18o_sd"))
                for iso in ("d13c","d18o"):
                    value=row["isotopes"].get(iso,{})
                    budget=value.get("budget",{})
                    item.update({f"{iso}_{key}":number for key,number in {
                        "raw":row[iso],"value":value.get("value"),"u_prec":value.get("u_prec"),"u_norm":value.get("u_norm"),
                        "u_corr":value.get("u_corr"),"u_combined":budget.get("u_combined"),"U":budget.get("expanded_uncertainty"),"k":budget.get("k")}.items()})
                whole_rows.append(item)
                if row["role"]=="unknown" and (command.include_outliers or not item["excluded"]):
                    rows.append(item)
        rows.sort(key=lambda row: (row["identifier1"], row["identifier2"], row["species"], row.get("acquired_at") or "", str(row["analysis"])))
        if not rows and command.format == "xlsx" and command.output_type == "dataset":
            rows = list(whole_rows)
        if not rows:
            raise ValueError("No evaluated unknown samples in this session/group")
        buffer=io.StringIO(newline="")
        writer=csv.DictWriter(buffer,fieldnames=list(rows[0])); writer.writeheader()
        for row in rows:
            writer.writerow({key:("'"+v if isinstance(v,str) and v.lstrip().startswith(("=","+","-","@")) else v) for key,v in row.items()})
        csv_bytes=buffer.getvalue().encode("utf-8-sig")
        # Group-filtered records exclude other sample groups even in the JSON dossier.
        allowed={r["evaluation_id"]+":"+r["analysis"] for r in rows}
        calculations=[{"evaluation_id":r["evaluation"]["id"],"normalization":r["evaluation"]["normalization"],
            "results":[v for v in r["evaluation"]["results"] if r["evaluation"]["id"]+":"+v["source_index"] in allowed]}
            for r in detail["runs"] if r["evaluation"]]
        payload={"simulation":self.repo.demo,"exported_at":self.repo.timestamp(),"review":command.model_dump(),
                 "session":{k:detail[k] for k in ("id","name","client","project","method_name","intended_use","qualification_id","processing_evidence")},
                 "method":detail["method"],"qualification":detail["qualification"],"results":rows,"calculations":calculations,
                 "correction_verification":[{"run_id":r["id"],"review":r["evaluation"].get("correction_review")} for r in detail["runs"] if r["evaluation"]],
                 "raw_source_inventory":detail["sources"]}
        json_bytes=json.dumps(payload,ensure_ascii=False,indent=2,allow_nan=False).encode("utf-8")
        content=csv_bytes if command.format=="csv" else json_bytes
        client_name = command.client_name if command.client_name is not None else detail["client"]
        series_name = command.series_name or command.group or detail["project"] or detail["name"]
        # The same series compaction, date and safe filename convention as IRMS processing.
        stem = _build_client_filename(client_name, pd.DataFrame({"Identifier": [series_name]}))[:-5]
        stem = _sanitize_filename(stem)[:180].rstrip(" .")
        pdf_bytes = build_session_dossier(payload) if command.format in ("pdf", "zip") else None
        xlsx_bytes = None
        if command.format in ("xlsx", "zip"):
            book = Workbook(); sheet = book.active; sheet.title = "Results"
            sheet.append(list(rows[0]))
            for row in rows:
                sheet.append([("'"+v if isinstance(v,str) and v.lstrip().startswith(("=","+","-","@")) else v) for v in row.values()])
            sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="17324D")
                sheet.column_dimensions[cell.column_letter].width = min(40, max(16, len(str(cell.value))+2))
            for record in sheet.iter_rows(min_row=2):
                for cell in record:
                    cell.alignment=Alignment(vertical="top")
                    if isinstance(cell.value,float): cell.number_format="0.0000"
            source_frame = pd.DataFrame({"Raw Label": [r["sample"] for r in rows], "Raw Comment": [r["sample_identifier"] for r in rows],
                "d 13C/12C  Mean": [r["d13c_raw"] for r in rows], "d 18O/16O  Mean": [r["d18o_raw"] for r in rows],
                "d 13C/12C  Std Dev": [r["d13c_internal_sd"] for r in rows], "d 18O/16O  Std Dev": [r["d18o_internal_sd"] for r in rows],
                "d13C_calibrated": [r["d13c_value"] for r in rows], "d18O_calibrated": [r["d18o_value"] for r in rows]})
            client_frame = _build_client_output_frame(source_frame, identifier_source=command.identifier_source,
                sample_source=command.sample_source, species_source="raw_label")
            client_frame = client_frame.drop(columns=[c for c in client_frame if c.startswith("__") or c == "Species"])
            for field, title in (("identifier1", "Identifier 1"), ("identifier2", "Identifier 2"), ("species", "Species")):
                client_frame[title] = [r[field] for r in rows]
            for iso in ("d13c", "d18o"):
                client_frame[f"{iso} expanded uncertainty / per mille"] = [r[f"{iso}_U"] for r in rows]
                client_frame[f"{iso} coverage factor k"] = [r[f"{iso}_k"] for r in rows]
            client_frame["Decision"] = [r["decision"] for r in rows]
            client_frame["Evaluation ID"] = [r["evaluation_id"] for r in rows]
            client_frame["Analysis"] = [r["analysis"] for r in rows]
            client_sheet = book.create_sheet("Client Output"); client_sheet.append(list(client_frame.columns))
            for values in client_frame.itertuples(index=False, name=None):
                client_sheet.append([None if pd.isna(v) else "'"+v if isinstance(v,str) and v.lstrip().startswith(("=","+","-","@")) else v for v in values])
            client_sheet.freeze_panes = "A2"; client_sheet.auto_filter.ref=client_sheet.dimensions
            for cell in client_sheet[1]:
                cell.font=Font(bold=True,color="FFFFFF");cell.fill=PatternFill("solid",fgColor="17324D")
                cell.alignment=Alignment(wrap_text=True,vertical="center");client_sheet.column_dimensions[cell.column_letter].width=22
            client_sheet.row_dimensions[1].height=40
            for values in client_sheet.iter_rows(min_row=2):
                for cell in values:
                    if isinstance(cell.value,float):cell.number_format="0.0000"
            buffer_xlsx=io.BytesIO(); book.save(buffer_xlsx); xlsx_bytes=buffer_xlsx.getvalue()
            # Reuse the original IRMS export writers, including numeric/species
            # formatting, duplicate highlighting and the separate outlier sheet.
            safe=lambda v: "'"+v if isinstance(v,str) and v.lstrip().startswith(("=","+","-","@")) else v
            clean_client=client_frame.map(safe)
            clean_client["__identifier_2_key"]=[safe(r["identifier2"]) for r in rows]
            stats=analysis["qc_statistics"]
            precision=(stats["d13c"]["final"]["sd"] or 0,stats["d18o"]["final"]["sd"] or 0,
                       stats["d13c"]["final"]["n"],stats["d18o"]["final"]["n"])
            client_bytes,_=build_client_output_workbook_bytes(source_frame,client_name=client_name,
                client_output_df=clean_client,precision_override=precision)
            all_frame=pd.DataFrame(whole_rows).map(safe)
            outlier_frame=all_frame.loc[all_frame["excluded"]].copy()
            retained=all_frame if command.include_outliers else all_frame.loc[~all_frame["excluded"]]
            dataset_bytes,_=build_dataset_workbook_bytes(retained,outliers=outlier_frame,client_name=client_name)
            if command.output_type=="client_output": xlsx_bytes=client_bytes
            elif command.output_type=="dataset": xlsx_bytes=dataset_bytes
        if command.format=="pdf": content=pdf_bytes
        if command.format=="xlsx": content=xlsx_bytes
        if command.format=="zip":
            archive=io.BytesIO()
            with zipfile.ZipFile(archive,"w",zipfile.ZIP_DEFLATED) as z:
                z.writestr("results.csv",csv_bytes); z.writestr("calculation-dossier.json",json_bytes)
                z.writestr(f"{stem}.xlsx", xlsx_bytes)
                z.writestr("client-output.xlsx", client_bytes)
                z.writestr("whole-results.xlsx", dataset_bytes)
                z.writestr("calculation-certificate.pdf", pdf_bytes)
                z.writestr("README.txt","IRMS Metrology Station\n"+("SIMULATED QUALIFICATION: observations may be real; see data_origin and calibration_verification per row.\n" if self.repo.demo else "")+"Decision is recorded per result. Review/blocked rows are not released results. All isotope values and uncertainties are in per mille VPDB. u_prec is individual QC SD; u_norm is sample-specific; u_corr propagates the shared correction coefficient. U=k*u_combined.\n")
            content=archive.getvalue()
        if command.format == "xlsx" and command.output_type == "dataset":
            stem += " - all data"
        sha=self.repo.blob(content)
        with self.repo.connect(write=True) as db:
            record=self.repo.insert(db,"session_exports",{"session_id":session_id,"sha256":sha,"filename":f"{stem}.{command.format}","format":command.format,"group":command.group,"rows":len(retained) if command.format=="xlsx" and command.output_type=="dataset" else len(rows),"output_type":command.output_type,"simulation":self.repo.demo})
            self.repo.audit(db,"session_results_exported",session_id,command.actor,command.reason,after=record)
        return record
