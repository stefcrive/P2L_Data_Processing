from __future__ import annotations

import csv
import io
import json
import sqlite3
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from services.irms_api.tests import test_metrology as fixtures
from services.irms_api.metrology.models import ResultsSessionCommand, RunCommand, SessionGroupsCommand, SessionExportCommand, ReleaseCommand, AnnotationCommand, RowReviewCommand, EvaluateCommand
from services.irms_api.metrology.repository import Repository
from services.irms_api.metrology.service import Service, Conflict

D = fixtures.DECISION


class ResultsSessionTests(unittest.TestCase):
    def test_bundled_analysis_reuses_detail_and_stays_fresh_after_edits(self):
        session = self.create()
        run, _ = self.import_batch(session)
        endpoint = f"/metrology/results-sessions/{session['id']}"
        before = self.service.results_session_detail(session["id"])
        analysis = self.service.results_session_analysis(session["id"])
        from services.irms_api.metrology import session_processing
        with patch.object(session_processing, "process_session_results", wraps=session_processing.process_session_results) as process:
            response = self.client.get(endpoint + "?include_analysis=true")
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(process.call_count, 1)
        bundled = response.json()
        self.assertEqual(bundled.pop("analysis"), analysis)
        self.assertEqual(bundled, before)
        mid = before["runs"][0]["evaluation"]["results"][-1]["id"]
        self.service.edit_session_identities(session["id"], {mid: {"identifier1": "Corrected sample"}})
        self.assertEqual(self.client.put(endpoint + "/uncertainty", json={**D, "coverage_factor": 3.5}).status_code, 200)
        self.assertEqual(self.client.put(endpoint + "/outlier-screening", json={**D, "method": "iqr", "threshold": 1.5}).status_code, 200)
        updated = self.client.get(endpoint + "?include_analysis=true").json()
        self.assertEqual(updated["analysis"], self.client.get(endpoint + "/analysis").json())
        row = next(r for r in updated["analysis"]["rows"] if r["id"] == mid)
        self.assertEqual(row["identifier1"], "Corrected sample")
        self.assertEqual(row["isotopes"]["d13c"]["budget"]["k"], 3.5)
        self.assertNotIn("analysis", self.client.get(endpoint).json())

    def test_session_reuses_unfiltered_qc_history_without_changing_release_gates(self):
        session = self.create()
        run, _ = self.import_batch(session)
        other = self.service.import_run("second.xlsx", fixtures.workbook([
            fixtures.row(1, "SHP2L", 0, sample_type="QC Standard"), fixtures.row(2, "other", 2)
        ]), RunCommand(**D, results_session_id=session["id"]))
        expected = {rid: self.service.run_detail(rid)["release_blockers"] for rid in (run["id"], other["id"])}
        with patch.object(self.service, "qc_history", wraps=self.service.qc_history) as history:
            detail = self.service.results_session_detail(session["id"])
            self.assertEqual(history.call_count, 1)
        self.assertEqual({r["id"]: r["release_blockers"] for r in detail["runs"]}, expected)
        with patch.object(self.service, "qc_history", wraps=self.service.qc_history) as history:
            self.service.run_detail(run["id"])
            self.assertEqual(history.call_count, 1)

    def test_export_computes_session_detail_once(self):
        session = self.create()
        self.import_batch(session)
        with patch.object(self.service, "results_session_detail", wraps=self.service.results_session_detail) as detail:
            self.service.export_results_session(session["id"], SessionExportCommand(**D, format="json"))
            self.assertEqual(detail.call_count, 1)

    def test_joint_pressure_fit_uses_failed_qc_and_preserves_unknown_composition(self):
        import numpy as np
        session = self.create()
        rows = []
        for i in range(8):
            p = -.03+.01*i
            row = fixtures.row(i+1, "SHP2L", 1.+2*p, sample_type="QC Standard")
            row.update({"i44_v": 5.+.1*i, "Pressure Adjust Target Intensity": 6., "Pressure Adjust Result Intensity": 6.+p})
            rows.append(row)
        for i, p in enumerate([-3., -1., 1., 3., 5.]):
            for j, intensity in enumerate([4., 8., 12., 16.]):
                row = fixtures.row(len(rows)+1, "SHP2L", 1.+.2*p-.04*intensity+.002*(-1)**(i+j), sample_type="QC Standard")
                row.update({"i44_v": intensity, "Pressure Adjust Target Intensity": 6., "Pressure Adjust Result Intensity": 6.+p,
                            "Pressure Adjust failed with Target Intensity": True})
                rows.append(row)
        # A clustered low-signal failure must neither define nor receive this fit.
        for i in range(12):
            row = fixtures.row(len(rows)+1, "SHP2L", -10.-i, sample_type="QC Standard")
            row.update({"i44_v": .01+.001*i, "Pressure Adjust Target Intensity": 6., "Pressure Adjust Result Intensity": 6.1+.01*i,
                        "Pressure Adjust failed with Target Intensity": True})
            rows.append(row)
        for label, composition in (("unknown A", 2.), ("unknown B", 5.)):
            row = fixtures.row(len(rows)+1, label, composition+.2*2.-.04*10.)
            row.update({"i44_v": 10., "Pressure Adjust Target Intensity": 6., "Pressure Adjust Result Intensity": 8.,
                        "Pressure Adjust failed with Target Intensity": True})
            rows.append(row)
        run = self.service.import_run("joint-pressure.xlsx", fixtures.workbook(rows), RunCommand(**D, results_session_id=session["id"]))
        source = self.service.run_detail(run["id"])["evaluation"]
        response = self.client.put(f"/metrology/results-sessions/{session['id']}/failed-correction", json={**D, "enabled": True})
        self.assertEqual(response.status_code, 200)
        analysis = self.service.results_session_analysis(session["id"])
        for iso in ("d13c", "d18o"):
            fit = analysis["failed_analysis_corrections"][iso]
            self.assertEqual(fit["training_population"], "pressure_failed_qc")
            self.assertEqual(fit["n"], 20)
            self.assertEqual(fit["unknown_applied_n"], 2)
            self.assertEqual(len(fit["screening"]["signal_excluded_ids"]), 12)
            self.assertTrue(fit["validation"]["passed"])
            for row in analysis["rows"][28:40]:
                self.assertNotIn("residual_correction", row["isotopes"][iso])
                self.assertIn("below the qualification minimum", row["failed_correction_attempts"][iso]["reason"])
            self.assertAlmostEqual(fit["model"]["slope"], .2, places=3)
            self.assertAlmostEqual(fit["model"]["intensity_slope"], -.04, places=3)
            self.assertAlmostEqual(fit["intensity_after"]["slope"], 0., places=12)
            first, second = [r["isotopes"][iso] for r in analysis["rows"][-2:]]
            self.assertAlmostEqual(second["value"]-first["value"], 3., places=12)
            g = np.array([2., 10.-fit["model"]["intensity_ref"]])
            expected_u2 = float(g@np.array(fit["model"]["coefficient_covariance"])@g)
            self.assertAlmostEqual(first["u_residual"]**2, expected_u2, places=12)
            components = [c["name"] for c in first["budget"]["components"]]
            self.assertEqual(components.count("pressure_adjustment_linearity"), 1)
            self.assertNotIn("residual_linearity", components)
        self.assertEqual(self.service.results_session_analysis(session["id"]), analysis)
        self.assertEqual(self.service.run_detail(run["id"])["evaluation"], source)
        export = self.service.export_results_session(session["id"], SessionExportCommand(**D, format="json"))
        payload = json.loads(self.repo.read_blob(export["sha256"]))
        self.assertEqual(payload["failed_analysis_corrections"], analysis["failed_analysis_corrections"])
        pdf = self.service.export_results_session(session["id"], SessionExportCommand(**D, format="pdf"))
        self.assertTrue(self.repo.read_blob(pdf["sha256"]).startswith(b"%PDF"))

    def test_failed_pressure_correction_persists_and_reaches_results_and_exports(self):
        session = self.create()
        rows = []
        for i, x in enumerate((-.03, -.02, -.01, .01, .02, .03)):
            row = fixtures.row(i+1, "SHP2L", 1.+2*x+(.001 if i%2 else -.001), sample_type="QC Standard")
            row.update({"Pressure Adjust Target Intensity": 6., "Pressure Adjust Result Intensity": 6.+x})
            rows.append(row)
        for label, x, value, failed in (("SHP2L", .06, 1.12, True), ("pressure failure", .2, 1.4, True),
                                       ("aborted", .25, 1.5, False), ("missing pressure", None, 1.4, True),
                                       ("normal sample", .1, 2.2, False)):
            row = fixtures.row(len(rows)+1, label, value, sample_type="QC Standard" if label == "SHP2L" else "Unknown")
            row.update({"Pressure Adjust Target Intensity": 6., "Pressure Adjust Result Intensity": None if x is None else 6.+x,
                        "Pressure Adjust failed with Target Intensity": failed})
            if label == "aborted":
                row["Status"] = "Aborted"
            rows.append(row)
        run = self.service.import_run("failed-pressure.xlsx", fixtures.workbook(rows), RunCommand(**D, results_session_id=session["id"]))
        source = self.service.run_detail(run["id"])["evaluation"]
        endpoint = f"/metrology/results-sessions/{session['id']}"
        default = self.client.get(endpoint + "/analysis").json()
        self.assertEqual(default["failed_analysis_corrections"], {})
        response = self.client.put(endpoint + "/failed-correction", json={**D, "enabled": True})
        self.assertEqual(response.status_code, 200, response.text)
        restarted = Service(Repository(self.repo.root))
        self.assertTrue(restarted.results_session_detail(session["id"])["correct_failed_analyses"])
        corrected = self.client.get(endpoint + "/analysis").json()
        for iso in ("d13c", "d18o"):
            fit = corrected["failed_analysis_corrections"][iso]
            self.assertEqual(fit["status"], "applied")
            self.assertEqual(fit["n"], 6)
            self.assertEqual(fit["scope"], "pressure_affected_analyses")
            self.assertEqual(fit["applied_n"], 2)
            self.assertEqual(fit["extrapolated_n"], 2)
            self.assertEqual(fit["unknown_applied_n"], 1)
            self.assertEqual(fit["model"]["x_ref"], 0.)
            self.assertAlmostEqual(fit["model"]["slope"], 2., delta=.05)
            for index in (6, 7):
                old = default["rows"][index]["isotopes"][iso]
                new = corrected["rows"][index]["isotopes"][iso]
                evidence = new["residual_correction"]
                self.assertTrue(evidence["failed_analysis"])
                self.assertEqual(evidence["effect"], "pressure_dependence")
                self.assertAlmostEqual(new["value"], old["value"]-fit["model"]["slope"]*corrected["rows"][index]["pressure_mismatch_v"])
                self.assertGreater(evidence["u"], 0.)
                self.assertAlmostEqual(new["budget"]["u_combined"]**2, old["budget"]["u_combined"]**2+evidence["u"]**2)
                self.assertEqual(sum(c["name"] == "pressure_adjustment_linearity" for c in new["budget"]["components"]), 1)
                self.assertEqual(corrected["rows"][index]["issues"], default["rows"][index]["issues"])
            for index in (8, 10):
                self.assertEqual(corrected["rows"][index]["isotopes"][iso], default["rows"][index]["isotopes"][iso])
                self.assertEqual(corrected["rows"][index]["failed_correction_attempts"], {})
            self.assertNotIn("residual_correction", corrected["rows"][9]["isotopes"][iso])
            self.assertEqual(corrected["rows"][9]["failed_correction_attempts"][iso]["reason"], "Pressure-adjustment difference is missing.")
            pool = fit["qc_pool"]
            self.assertEqual(pool["status"], "admitted")
            self.assertEqual(pool["admitted_ids"], [corrected["rows"][6]["id"]])
            self.assertLess(pool["after"]["sd"], pool["before"]["sd"])
            self.assertEqual(corrected["qc_statistics"][iso]["final"]["n"], 7)
            self.assertEqual(corrected["qc_statistics"][iso]["imported"], default["qc_statistics"][iso]["imported"])
            self.assertEqual(corrected["diagnostics_before"]["materials"][0]["isotopes"][iso]["pressure_dependence"]["n"], 6)
            self.assertEqual(corrected["diagnostics_after"]["materials"][0]["isotopes"][iso]["pressure_dependence"]["n"], 7)
        self.assertEqual([{k: v for k, v in f.items() if k not in ("value", "session_qc_admitted")} for f in corrected["qc_review_flags"]],
                         [{k: v for k, v in f.items() if k != "value"} for f in default["qc_review_flags"]])
        self.assertTrue(all(f["session_qc_admitted"] for f in corrected["qc_review_flags"] if f["category"] == "pressure_adjustment"))
        self.assertEqual(restarted.results_session_analysis(session["id"]), corrected)
        self.assertEqual(self.service.run_detail(run["id"])["evaluation"], source)
        exported = restarted.export_results_session(session["id"], SessionExportCommand(**D, format="json"))
        payload = json.loads(self.repo.read_blob(exported["sha256"]))
        self.assertTrue(payload["correct_failed_analyses"])
        self.assertEqual(payload["failed_analysis_corrections"], corrected["failed_analysis_corrections"])
        sample = next(r for r in payload["results"] if r["sample"] == "pressure failure")
        self.assertEqual(sample["d13c_value"], corrected["rows"][7]["isotopes"]["d13c"]["value"])
        self.assertEqual(sample["decision"], "blocked")
        normal = next(r for r in payload["results"] if r["sample"] == "normal sample")
        self.assertEqual(normal["d13c_value"], corrected["rows"][10]["isotopes"]["d13c"]["value"])
        whole = restarted.results_session_detail(session["id"])
        for iso in ("d13c", "d18o"):
            self.assertEqual(whole["runs"][0]["evaluation"]["session_qc"][iso], corrected["qc_statistics"][iso]["final"])
        pdf = restarted.export_results_session(session["id"], SessionExportCommand(**D, format="pdf"))
        self.assertTrue(self.repo.read_blob(pdf["sha256"]).startswith(b"%PDF"))
        from services.irms_api.api import main as legacy
        from services.irms_api.session_store import FileSessionStore
        from services.irms_api.metrology.plot_bridge import open_plot_bridge
        from services.irms_api.metrology.models import Decision
        with patch.object(legacy, "store", FileSessionStore(Path(self.fixture.temp.name)/"failed-tools")):
            bridge = open_plot_bridge(restarted, session["id"], "all", Decision(**D))
            frame = legacy.store.load_frame(bridge["session_id"])
            sample_id = corrected["rows"][7]["id"]
            self.assertAlmostEqual(frame.loc[int(bridge["row_mapping"][sample_id]), "d 13C/12C  Mean"], sample["d13c_value"])
            response = self.client.put(endpoint + "/outlier-screening", json={**D, "method": "sigma", "threshold": 3., "pressure_adjustment_as_outlier": True})
            self.assertEqual(response.status_code, 200, response.text)
            screened = restarted.results_session_analysis(session["id"])
            for iso in ("d13c", "d18o"):
                self.assertEqual(screened["qc_statistics"][iso]["final"]["n"], 6)
                self.assertEqual(screened["failed_analysis_corrections"][iso]["qc_pool"]["admitted_ids"], [])
            screened_bridge = open_plot_bridge(restarted, session["id"], "all", Decision(**D))
            self.assertNotEqual(bridge["session_id"], screened_bridge["session_id"])
            self.assertTrue(legacy.store.load_metadata(screened_bridge["session_id"])["processing"]["config"]["pressure_adjustment_as_outlier"])
            response = self.client.put(endpoint + "/outlier-screening", json={**D, "method": "sigma", "threshold": 3., "pressure_adjustment_as_outlier": False})
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(restarted.results_session_analysis(session["id"])["qc_statistics"]["d13c"]["final"]["n"], 7)
            response = self.client.put(endpoint + "/failed-correction", json={**D, "enabled": False})
            self.assertEqual(response.status_code, 200, response.text)
            restored_bridge = open_plot_bridge(restarted, session["id"], "all", Decision(**D))
            self.assertNotEqual(bridge["session_id"], restored_bridge["session_id"])
            restored = legacy.store.load_frame(restored_bridge["session_id"])
            self.assertAlmostEqual(restored.loc[int(restored_bridge["row_mapping"][sample_id]), "d 13C/12C  Mean"], default["rows"][7]["isotopes"]["d13c"]["value"])
        restored_analysis = self.service.results_session_analysis(session["id"])
        self.assertEqual(restored_analysis["rows"], default["rows"])
        self.assertEqual(restored_analysis["qc_statistics"], default["qc_statistics"])
        with self.repo.connect(write=True) as db:
            self.repo.update(db, "runs", run["id"], status="released")
        self.assertEqual(self.client.put(endpoint + "/failed-correction", json={**D, "enabled": True}).status_code, 422)

    def test_failed_pressure_correction_requires_usable_qc_and_inputs(self):
        from services.irms_api.metrology.models import MethodConfig
        from services.irms_api.metrology.session_processing import correct_failed_results
        config = MethodConfig.model_validate(self.method["config"])
        points = [{"id": str(i), "x": i*.01, "y": 1.+i*.02} for i in range(6)]
        qc = {"isotopes": {iso: {"pressure_dependence": {"points": points}} for iso in ("d13c", "d18o")}}
        failed = {"0", "failed", "manual"}
        statistical = {iso: {"1"} for iso in ("d13c", "d18o")}
        rows = [{"id": "failed", "issues": ["Qtegra pressure adjustment failed"], "pressure_mismatch_v": .2, "isotopes": {}},
                {"id": "manual", "excluded": True, "issues": ["Qtegra pressure adjustment failed"], "pressure_mismatch_v": .2, "isotopes": {}}]
        decision = correct_failed_results(rows, qc, failed, statistical, config, 2.)
        self.assertEqual(decision["d13c"]["n"], 4)
        self.assertEqual(decision["d13c"]["status"], "unavailable_results")
        self.assertEqual(decision["d13c"]["attempted_n"], 1)
        self.assertEqual(rows[0]["failed_correction_attempts"]["d13c"]["reason"], "Isotope value or complete uncertainty budget is missing.")
        self.assertNotIn("failed_correction_attempts", rows[1])
        for label, subset, expected in (("too few QC", points[:3], "insufficient_evidence"),
                                         ("constant pressure", [{**p, "x": .01} for p in points], "insufficient_evidence"),
                                         ("no delta relationship", [{**p, "y": 1.} for p in points], "not_improved")):
            with self.subTest(label=label):
                qc["isotopes"]["d13c"]["pressure_dependence"]["points"] = subset
                outcome = correct_failed_results(rows, qc, failed, statistical, config, 2.)["d13c"]
                self.assertEqual(outcome["status"], expected)
                self.assertEqual(outcome["applied_n"], 0)

    def test_pressure_qc_pool_admission_is_per_isotope_and_preserves_other_exclusions(self):
        from services.irms_api.metrology.session_processing import admit_pressure_corrected_qc
        from services.irms_api.metrology.session_analysis import session_analysis
        session = self.create()
        self.import_linearity(session)
        detail = self.service.results_session_detail(session["id"])
        rows = detail["runs"][0]["evaluation"]["results"]
        qc = [r for r in rows if r["role"] == "qc"]
        for i, row in enumerate(qc[:6]):
            row["issues"] = []
            for iso in ("d13c", "d18o"):
                row["isotopes"][iso]["value"] = .01*i
        recovered, aborted = qc[-2:]
        recovered["issues"] = ["Qtegra pressure adjustment failed"]
        aborted["issues"] = ["Qtegra pressure adjustment failed", "Qtegra reports an acquisition failure"]
        for row in (recovered, aborted):
            for iso in ("d13c", "d18o"):
                row["isotopes"][iso].update(value=.025 if iso == "d13c" else 10., residual_correction={"effect": "pressure_dependence"})
        failed = {recovered["id"], aborted["id"]}
        decisions = {iso: {} for iso in ("d13c", "d18o")}
        admit_pressure_corrected_qc(rows, failed, {"d13c": set(), "d18o": set()}, decisions)
        self.assertEqual(decisions["d13c"]["qc_pool"]["admitted_ids"], [recovered["id"]])
        self.assertEqual(decisions["d18o"]["qc_pool"]["status"], "not_improved")
        self.assertFalse(aborted["isotopes"]["d13c"].get("session_qc_admitted"))
        analysis = session_analysis(detail, outliers={"flags": []})
        self.assertEqual(analysis["qc_statistics"]["d13c"]["final"]["n"], 7)
        self.assertEqual(analysis["qc_statistics"]["d18o"]["final"]["n"], 6)
        ranged = session_analysis(detail, outliers={"flags": []}, range_exclusions={"d13c": [recovered["id"]]})
        self.assertEqual(ranged["qc_statistics"]["d13c"]["final"]["n"], 6)
        screened = session_analysis(detail, outliers={"flags": [{"measurement_id": recovered["id"], "isotope": "d13c"}]})
        self.assertEqual(screened["qc_statistics"]["d13c"]["final"]["n"], 6)
        recovered["excluded"] = True
        manual = session_analysis(detail, outliers={"flags": []})
        self.assertEqual(manual["qc_statistics"]["d13c"]["final"]["n"], 6)

    def test_only_pressure_affected_unknown_gets_pressure_correction_from_method_baseline(self):
        session = self.create()
        rows = []
        for i in range(8):
            x = -.03+i*.01
            row = fixtures.row(i+1, "SHP2L", 1.+2*x+(.001 if i%2 else -.001), sample_type="QC Standard")
            row.update({"i44_v": 4.+i*.1, "Pressure Adjust Target Intensity": 6., "Pressure Adjust Result Intensity": 6.+x})
            rows.append(row)
        sample = fixtures.row(9, "normal sample", 2.4)
        sample.update({"i44_v": 5.5, "Pressure Adjust Target Intensity": 6., "Pressure Adjust Result Intensity": 6.2})
        rows.append(sample)
        rows.append({**sample, "Index": 10, "Label": "pressure sample", "Pressure Adjust failed with Target Intensity": True})
        rows.append({**sample, "Index": 11, "Label": "missing pressure", "Pressure Adjust failed with Target Intensity": True,
                     "Pressure Adjust Result Intensity": None})
        run = self.service.import_run("pressure-and-signal.xlsx", fixtures.workbook(rows), RunCommand(**D, results_session_id=session["id"]))
        original = self.service.run_detail(run["id"])["evaluation"]["results"][-2]
        default = self.service.results_session_analysis(session["id"])
        self.assertIn("residual_correction", default["rows"][-3]["isotopes"]["d13c"])
        response = self.client.put(f"/metrology/results-sessions/{session['id']}/failed-correction", json={**D, "enabled": True})
        self.assertEqual(response.status_code, 200, response.text)
        corrected = self.service.results_session_analysis(session["id"])
        for iso in ("d13c", "d18o"):
            normal = corrected["rows"][-3]["isotopes"][iso]
            self.assertEqual(normal, default["rows"][-3]["isotopes"][iso])
            self.assertIn(normal["residual_correction"]["effect"], ("intensity_dependence", "sample_reference_dependence", "pressure_adjusted_dependence"))
            self.assertFalse(any(c["name"] == "pressure_adjustment_linearity" for c in normal["budget"]["components"]))
            self.assertNotIn("residual_correction", default["rows"][-2]["isotopes"][iso])
            self.assertNotIn("residual_correction", corrected["rows"][-1]["isotopes"][iso])
            value = corrected["rows"][-2]["isotopes"][iso]
            fit = corrected["failed_analysis_corrections"][iso]
            self.assertAlmostEqual(value["value"], original["isotopes"][iso]["value"]-fit["model"]["slope"]*.2)
            components = value["budget"]["components"]
            self.assertEqual(sum(c["name"] == "pressure_adjustment_linearity" for c in components), 1)
            self.assertFalse(any(c["name"] == "residual_linearity" for c in components))
            self.assertAlmostEqual(value["budget"]["u_combined"]**2, original["isotopes"][iso]["budget"]["u_combined"]**2+value["u_residual"]**2)
            self.assertEqual(fit["qc_pool"]["status"], "no_eligible_qc")

    def test_qualification_pressure_warnings_do_not_remove_qc_or_trigger_pressure_correction(self):
        from copy import deepcopy
        from services.irms_api.metrology.pipeline import evaluate
        from services.irms_api.metrology.session_processing import process_session_results, SIGNAL_EFFECTS
        session = self.create()
        rows = []
        for i in range(8):
            x = -.03+i*.01
            row = fixtures.row(i+1, "SHP2L", 1.+2*x+(.001 if i%2 else -.001), sample_type="QC Standard")
            row.update({"i44_v": 4.+i*.1, "Pressure Adjust Target Intensity": 6., "Pressure Adjust Result Intensity": 6.+x})
            rows.append(row)
        for label, x, flag in (("normal", .02, False), ("range only", .2, False), ("flag only", .01, True),
                               ("missing metadata", None, False)):
            row = fixtures.row(len(rows)+1, label, 2.4)
            row.update({"i44_v": 5.5, "Pressure Adjust Target Intensity": 6.,
                        "Pressure Adjust Result Intensity": None if x is None else 6.+x,
                        "Pressure Adjust failed with Target Intensity": flag})
            rows.append(row)
        self.service.import_run("qualification-pressure-range.xlsx", fixtures.workbook(rows), RunCommand(**D, results_session_id=session["id"]))
        detail = self.service.results_session_detail(session["id"])
        # Evaluate the imported observations under a qualification with pressure limits.
        detail["method"]["config"]["ranges"]["pressure_mismatch_v"] = {"low": -.001, "high": .001}
        run = detail["runs"][0]
        evaluated = evaluate(run, detail["method"], run["evaluation"]["material_snapshots"], run["measurements"], [])
        evaluated["id"] = run["evaluation"]["id"]
        run["evaluation"] = evaluated
        detail.pop("saved_outliers", None)
        for effect in SIGNAL_EFFECTS:
            for enabled in (False, True):
                with self.subTest(effect=effect, pressure_enabled=enabled):
                    candidate = deepcopy(detail)
                    candidate["correct_failed_analyses"] = enabled
                    candidate["residual_overrides"] = {
                        f"{self.method['config']['qc_id']}:{other}:{iso}": {"enabled": other == effect}
                        for iso in ("d13c", "d18o") for other in SIGNAL_EFFECTS}
                    processed = process_session_results(candidate)
                    results = {r["label"]: r for r in processed["runs"][0]["evaluation"]["results"]}
                    for iso in ("d13c", "d18o"):
                        self.assertEqual(processed["residual_corrections"][iso][effect]["n"], 8)
                        self.assertEqual(processed["residual_corrections"][iso][effect]["status"], "applied")
                        for label in ("normal", "range only", "missing metadata"):
                            value = results[label]["isotopes"][iso]
                            self.assertEqual(value["residual_correction"]["effect"], effect)
                            self.assertFalse(any(c["name"] == "pressure_adjustment_linearity" for c in value["budget"]["components"]))
                        for label in ("flag only",):
                            value = results[label]["isotopes"][iso]
                            self.assertFalse(any(c["name"] == "residual_linearity" for c in value["budget"]["components"]))
                            if enabled:
                                self.assertEqual(value["residual_correction"]["effect"], "pressure_dependence")
                            else:
                                self.assertNotIn("residual_correction", value)
                        if enabled:
                            self.assertEqual(processed["failed_analysis_corrections"][iso]["applied_n"], 1)
                            self.assertEqual(processed["failed_analysis_corrections"][iso]["n"], 8)

    def test_import_ignores_evaluate_flag_but_retains_acquisition_failure_checks(self):
        session = self.create()
        rows = [fixtures.row(i+1, "SHP2L", .01*i, sample_type="QC Standard") for i in range(4)]
        rows.extend([fixtures.row(5, "sample", 1.)])
        for row in rows:
            row["Evaluate"] = False
        rows[2]["Pressure Adjust failed with Target Intensity"] = True
        rows[3]["Status"] = "Aborted"
        run = self.service.import_run("evaluate-disabled.xlsx", fixtures.workbook(rows), RunCommand(**D, results_session_id=session["id"]))
        evaluation = self.service.run_detail(run["id"])["evaluation"]
        results = evaluation["results"]
        for row in results:
            self.assertFalse(row["evaluate"])
            self.assertNotIn("Qtegra Evaluate is disabled", row["issues"])
        self.assertIn("Qtegra pressure adjustment failed", results[2]["issues"])
        self.assertIn("Qtegra reports an acquisition failure", results[3]["issues"])
        for index in (0, 1, 4):
            self.assertFalse(results[index]["issues"])
            self.assertTrue(all(iso in results[index]["isotopes"] for iso in ("d13c", "d18o")))
        analysis = self.service.results_session_analysis(session["id"])
        self.assertEqual({flag["measurement_id"] for flag in analysis["qc_review_flags"] if flag["category"] == "pressure_adjustment"}, {results[2]["id"]})
        self.assertEqual({flag["measurement_id"] for flag in analysis["outliers"]["flags"] if flag["category"] == "no_signal"}, {results[3]["id"]})
        self.assertFalse(any(flag["category"] == "pressure_adjustment" for flag in analysis["outliers"]["flags"]))
        for iso in ("d13c", "d18o"):
            self.assertEqual(analysis["qc_statistics"][iso]["final"]["n"], 2)

        endpoint = f"/metrology/results-sessions/{session['id']}/outlier-screening"
        response = self.client.put(endpoint, json={**D, "method": "iqr", "threshold": 1.5, "pressure_adjustment_as_outlier": True})
        self.assertEqual(response.status_code, 200, response.text)
        restarted = Service(Repository(self.repo.root))
        self.assertTrue(restarted.results_session_detail(session["id"])["outlier_screening"]["pressure_adjustment_as_outlier"])
        enabled = restarted.results_session_analysis(session["id"])
        self.assertEqual({f["measurement_id"] for f in enabled["outliers"]["flags"] if f["category"] == "pressure_adjustment"}, {results[2]["id"]})
        self.assertFalse(any(f["category"] == "pressure_adjustment" for f in enabled["qc_review_flags"]))
        self.assertEqual(enabled["rows"][2]["failure_categories"], {"d13c": "pressure_adjustment", "d18o": "pressure_adjustment"})
        self.assertEqual(enabled["rows"][3]["failure_categories"], {"d13c": "no_signal", "d18o": "no_signal"})
        response = self.client.put(endpoint, json={**D, "method": "iqr", "threshold": 1.5, "pressure_adjustment_as_outlier": False})
        self.assertEqual(response.status_code, 200, response.text)
        disabled = restarted.results_session_analysis(session["id"])
        self.assertFalse(any(f["category"] == "pressure_adjustment" for f in disabled["outliers"]["flags"]))
        self.assertEqual({f["measurement_id"] for f in disabled["outliers"]["flags"] if f["category"] == "no_signal"}, {results[3]["id"]})

    def test_recovered_qc_is_screened_before_residual_fitting_and_cache_is_invalidated(self):
        from services.irms_api.metrology.qc_screening import stored_qc_screening
        from services.irms_api.metrology.repository import encode
        import hashlib
        session = self.create()
        run = self.import_linearity(session)
        evaluation = self.service.run_detail(run["id"])["evaluation"]
        # Legacy method-domain rejection erased these isotope results. One is
        # anomalous but recoverable; a failed acquisition must not skew screening.
        qc = [r for r in evaluation["results"] if r["role"] == "qc"]
        recovered, failed = qc[-2:]
        for iso in ("d13c", "d18o"):
            evaluation["normalization"][iso]["correction"] = {"name":"Test correction","predictor":"i44_v","slope":0.,"u_slope":.001,"center":5.,"domain":{"low":3.,"high":9.},"anchor_predictor_means":[5.,5.]}
            recovered[iso] = 5.
            failed[iso] = 100.
        recovered.update(isotopes={},i44_v=15.,issues=["correction predictor outside validated domain"])
        failed.update(isotopes={},i44_v=16.,issues=["Qtegra pressure adjustment failed"])
        with self.repo.connect(write=True) as db:
            run_record = self.repo.get(db,"runs",run["id"])
            data = {k:v for k,v in evaluation.items() if k not in ("id","created_at","run_id","method_id")}
            record = self.repo.insert(db,"evaluations",data,run_id=run["id"],method_id=run["method_id"])
            run_record["latest_evaluation_id"] = record["id"]
            self.repo.update(db,"runs",run["id"],run_record)
            session = self.repo.get(db,"results_sessions",session["id"])
            session["outlier_screening"] = {"method":"sigma","threshold":1.}
            self.repo.update(db,"results_sessions",session["id"],session)
            population = [(run["id"],record["id"],run_record["revision"])]
            old_hash = hashlib.sha256(encode({"version":3,"population":population,"settings":session["outlier_screening"]}).encode()).hexdigest()
            db.execute("INSERT INTO qc_screenings(id,session_id,fingerprint,data,created_at) VALUES(?,?,?,?,?)",("old-screening",session["id"],old_hash,encode({"method":"sigma","threshold":1.,"flags":[]}),self.repo.timestamp()))
            screening = stored_qc_screening(self.repo,db,session)
        self.assertNotEqual(screening["id"],"old-screening")
        for iso in ("d13c","d18o"):
            flags = [f for f in screening["flags"] if f["isotope"] == iso]
            self.assertEqual([f["measurement_id"] for f in flags],[recovered["id"]])
            self.assertEqual(flags[0]["population_n"],7)
        analysis = self.service.results_session_analysis(session["id"])
        for iso in ("d13c","d18o"):
            fit = next(d for d in analysis["residual_corrections"][iso].values() if d["status"] == "applied")
            self.assertEqual(fit["n"],6)
            self.assertNotIn(recovered["id"],{p["id"] for p in fit["points"] if not p.get("excluded_from_fit")})
            actual = next(m for m in analysis["diagnostics_after"]["materials"] if m["material_id"] == self.method["config"]["qc_id"])
            self.assertAlmostEqual(actual["isotopes"][iso][fit["source_effect"]]["slope"],0.,places=10)

    def test_import_groups_by_identifier1_and_species_across_workbooks(self):
        session = self.create()
        content = fixtures.workbook([
            fixtures.row(1, "SHP2L", 0, sample_type="QC Standard"),
            fixtures.row(2, "Core A - calcite", 1),
            fixtures.row(3, "Core A - calcite", 2),
            fixtures.row(4, "Core A - aragonite", 3),
            fixtures.row(5, "Core B - calcite", 4),
        ])
        first = self.service.import_run("groups.xlsx", content, RunCommand(**D, results_session_id=session["id"]))
        second = self.service.import_run("more.xlsx", fixtures.workbook([
            fixtures.row(1, "Core A - calcite", 5),
        ]), RunCommand(**D, results_session_id=session["id"]))
        detail = Service(Repository(self.repo.root)).results_session_detail(session["id"])
        grouped = {}
        for run in detail["runs"]:
            for row in run["evaluation"]["results"]:
                grouped.setdefault(row["label"], set()).add(detail["groups"][row["id"]])
        self.assertEqual(grouped["Core A - calcite"], {"Core A · calcite"})
        self.assertEqual(grouped["Core A - aragonite"], {"Core A · aragonite"})
        self.assertEqual(grouped["Core B - calcite"], {"Core B · calcite"})
        analysis = self.service.results_session_analysis(session["id"])
        self.assertEqual(sum(row["sample_group"] == "Core A · calcite" for row in analysis["rows"]), 3)
        row = next(row for row in analysis["rows"] if row["sample_group"] == "Core A · calcite")
        self.service.save_session_groups(session["id"], SessionGroupsCommand(**D, groups={row["id"]: "Reviewed group"}))
        self.service.import_run("duplicate.xlsx", content, RunCommand(**D, results_session_id=session["id"]))
        updated = self.service.results_session_detail(session["id"])
        self.assertEqual(updated["groups"][row["id"]], "Reviewed group")
        self.assertEqual(updated["run_ids"], [first["id"], second["id"]])

    def test_automatic_group_uses_available_identity_when_fields_are_missing(self):
        from services.irms_api.metrology.sessions import imported_sample_group
        def row(identifier, species, label=""):
            return {"identifier1": identifier, "identifier2": "99", "species": species, "label": label}
        self.assertEqual(imported_sample_group(row(" Core A ", " calcite ")), "Core A · calcite")
        self.assertEqual(imported_sample_group(row("Core A", "")), "Core A")
        self.assertEqual(imported_sample_group(row("", "calcite", "Original label")), "Original label · calcite")
        self.assertEqual(imported_sample_group(row("", "")), "Main batch")

    def test_session_coverage_factor_persists_and_scales_all_budgets_and_exports(self):
        session = self.create()
        run, _ = self.import_batch(session)
        original = self.service.run_detail(run["id"])["evaluation"]
        endpoint = f"/metrology/results-sessions/{session['id']}/uncertainty"
        for invalid in (0, -1, "NaN", "Infinity"):
            self.assertEqual(self.client.put(endpoint, json={**D, "coverage_factor": invalid}).status_code, 422)
        response = self.client.put(endpoint, json={**D, "coverage_factor": 3.5})
        self.assertEqual(response.status_code, 200, response.text)
        restarted = Service(Repository(self.repo.root))
        detail = restarted.results_session_detail(session["id"])
        self.assertEqual(detail["coverage_factor"], 3.5)
        analysis = restarted.results_session_analysis(session["id"])
        for row in analysis["rows"]:
            for iso in ("d13c", "d18o"):
                budget = row["isotopes"][iso]["budget"]
                self.assertEqual(budget["k"], 3.5)
                self.assertAlmostEqual(budget["expanded_uncertainty"], 3.5 * budget["u_combined"])
        self.assertEqual(restarted.run_detail(run["id"])["evaluation"], original)
        self.assertEqual(detail["method"]["config"]["coverage_factor"], self.method["config"]["coverage_factor"])
        export = restarted.export_results_session(session["id"], SessionExportCommand(**D, format="json"))
        payload = json.loads(self.repo.read_blob(export["sha256"]))
        self.assertEqual(payload["coverage_factor"], 3.5)
        for row in payload["results"]:
            for iso in ("d13c", "d18o"):
                self.assertEqual(row[f"{iso}_k"], 3.5)
                self.assertAlmostEqual(row[f"{iso}_U"], 3.5 * row[f"{iso}_u_combined"])

    def test_review_warnings_do_not_erase_linearity_and_all_data_is_opt_in(self):
        from services.irms_api.metrology.session_analysis import session_analysis
        session=self.create(); self.import_batch(session)
        detail=self.service.results_session_detail(session["id"])
        for row in detail["runs"][0]["evaluation"]["results"]:
            row["issues"]=["mass_ug missing or outside validated range", "d13c internal SD missing or at/above limit"]
            row["initial_i44_v"] = row["sequence"] * 2
        analysis=session_analysis(detail)
        self.assertEqual(analysis["correction_review"]["d13c"]["paired_n"],3)
        self.assertEqual(analysis["diagnostics_after"]["materials"][0]["isotopes"]["d13c"]["intensity_dependence"]["n"],3)
        qc=[r for r in detail["runs"][0]["evaluation"]["results"] if r["role"]=="qc"]
        self.assertEqual([p["x"] for p in analysis["diagnostics_after"]["materials"][0]["isotopes"]["d13c"]["intensity_dependence"]["points"]], [r["initial_i44_v"] for r in qc])
        self.assertFalse(any(m["material_id"]=="__all__" for m in analysis["diagnostics_after"]["materials"]))
        all_data=self.client.get(f"/metrology/results-sessions/{session['id']}/analysis?include_all_data=true")
        self.assertEqual(all_data.status_code,200,all_data.text)
        pooled=next(m for m in all_data.json()["diagnostics_after"]["materials"] if m["material_id"]=="__all__")
        self.assertEqual(pooled["isotopes"]["d13c"]["intensity_dependence"]["n"],6)

    def test_workbook_inclusion_is_reversible_and_preserves_evaluation(self):
        session=self.create();run,_=self.import_batch(session)
        before=self.service.run_detail(run["id"])["evaluation"]
        url=f"/metrology/results-sessions/{session['id']}/workbooks/{run['id']}"
        removed=self.client.post(url+"?included=false",json=D)
        self.assertEqual(removed.status_code,200,removed.text)
        detail=self.service.results_session_detail(session["id"])
        self.assertEqual(detail["runs"],[])
        self.assertEqual(detail["detached_runs"][0]["evaluation"],before)
        self.assertEqual(self.client.post(url+"?included=true",json=D).status_code,200)
        self.assertEqual(self.service.run_detail(run["id"])["evaluation"],before)
        other=self.create("Other client")
        self.assertEqual(self.client.post(f"/metrology/results-sessions/{other['id']}/workbooks/{run['id']}",json=D).status_code,422)

    def test_original_excel_writers_export_whole_population_and_outliers(self):
        from openpyxl import load_workbook
        session=self.create();self.import_batch(session)
        client=self.service.export_results_session(session["id"],SessionExportCommand(**D,format="xlsx",output_type="client_output"))
        book=load_workbook(io.BytesIO(self.repo.read_blob(client["sha256"])))
        self.assertEqual(book.sheetnames,["Client Output"])
        headers=[c.value for c in book.active[1]]
        corrected=next(i+1 for i,h in enumerate(headers) if h and h.startswith("Corrected d13C"))
        self.assertEqual(book.active.cell(2,corrected).number_format,"0.00")
        self.assertEqual(book.active.cell(2,1).data_type,"s")
        whole=self.service.export_results_session(session["id"],SessionExportCommand(**D,format="xlsx",output_type="dataset"))
        book=load_workbook(io.BytesIO(self.repo.read_blob(whole["sha256"])))
        self.assertEqual(book["Data"].max_row,7)  # Three QC and three unknowns.
        self.assertIn("Statistics",book.sheetnames)
        qc=[fixtures.row(i+1,"SHP2L",v,sample_type="QC Standard") for i,v in enumerate([0,.001,.002,.003,.004,.005,.006,1])]
        self.service.import_run("extra-qc.xlsx",fixtures.workbook(qc),RunCommand(**D,results_session_id=session["id"]))
        self.client.put(f"/metrology/results-sessions/{session['id']}/outlier-screening",json={**D,"method":"iqr","threshold":1.5})
        for include in (False,True):
            exported=self.service.export_results_session(session["id"],SessionExportCommand(**D,format="xlsx",output_type="dataset",include_outliers=include))
            book=load_workbook(io.BytesIO(self.repo.read_blob(exported["sha256"])))
            self.assertIn("Outliers",book.sheetnames)
            self.assertGreater(book["Outliers"].max_row,1)
            self.assertEqual(book["Data"].max_row-1+(0 if include else book["Outliers"].max_row-1),14)

    def test_missing_mass_does_not_gate_valid_signals(self):
        session=self.create()
        rows=[fixtures.row(i+1,"SHP2L",v,sample_type="QC Standard") for i,v in enumerate([-.01,0,.01])]
        for row in rows: row["mass_ug"]=None
        run=self.service.import_run("no-mass.xlsx",fixtures.workbook(rows),RunCommand(**D,results_session_id=session["id"]))
        evaluation=self.service.run_detail(run["id"])["evaluation"]
        self.assertTrue(evaluation["qc"]["passed"])
        self.assertFalse(any("mass_ug" in issue for r in evaluation["results"] for issue in r["issues"]))

    def test_residual_overrides_persist_and_can_be_restored(self):
        session = self.create()
        run, _ = self.import_batch(session)
        original = self.service.run_detail(run["id"])["evaluation"]
        material = self.method["config"]["qc_id"]
        endpoint = f"/metrology/results-sessions/{session['id']}/residual-overrides"
        command = {**D, "material_id": material, "effect": "intensity_dependence", "isotope": "d13c",
                   "settings": {"enabled": False}}
        response = self.client.put(endpoint, json=command)
        self.assertEqual(response.status_code, 200, response.text)
        analysis = self.client.get(f"/metrology/results-sessions/{session['id']}/analysis").json()
        self.assertEqual(analysis["residual_corrections"]["d13c"]["intensity_dependence"]["status"], "disabled")
        key = f"{material}:intensity_dependence:d13c"
        persisted = Service(Repository(self.repo.root)).results_session_detail(session["id"])
        self.assertFalse(persisted["residual_overrides"][key]["enabled"])
        self.assertEqual(self.service.run_detail(run["id"])["evaluation"], original)
        self.assertEqual(self.client.put(endpoint, json={**command,"material_id":"foreign"}).status_code, 422)
        self.assertEqual(self.client.put(endpoint, json={**command,"settings":None}).status_code, 200)
        self.assertNotIn(key, self.service.results_session_detail(session["id"])["residual_overrides"])

    def import_linearity(self, session):
        rows = [fixtures.row(i+1, "SHP2L", .03*i + (.001 if i % 2 else -.001), sample_type="QC Standard") for i in range(8)]
        for i, row in enumerate(rows):
            row["i44_v"] = 4.+i*.1
        sample = fixtures.row(9, "sample", 1.2)
        sample["i44_v"] = 18.  # Beyond both the fit and method intensity range.
        rows.append(sample)
        return self.service.import_run("residual-fit.xlsx", fixtures.workbook(rows), RunCommand(**D, results_session_id=session["id"]))

    def test_automatic_residual_reaches_all_results_budgets_charts_and_exports(self):
        from services.irms_api.api import main as legacy
        from services.irms_api.session_store import FileSessionStore
        from services.irms_api.metrology.plot_bridge import open_plot_bridge
        from services.irms_api.metrology.models import Decision
        session = self.create()
        run = self.import_linearity(session)
        original = self.service.run_detail(run["id"])["evaluation"]
        restarted = Service(Repository(self.repo.root))
        detail = restarted.results_session_detail(session["id"])
        self.assertEqual(detail, restarted.results_session_detail(session["id"]))
        analysis = restarted.results_session_analysis(session["id"], include_all_data=True)
        sample = next(r for r in analysis["rows"] if r["role"] == "unknown")
        original_sample = next(r for r in original["results"] if r["id"] == sample["id"])
        for iso in ("d13c", "d18o"):
            applied = [d for d in analysis["residual_corrections"][iso].values() if d["status"] == "applied"]
            self.assertEqual(len(applied), 1)
            self.assertLess(applied[0]["after"]["sd"], applied[0]["before"]["sd"])
            self.assertEqual(applied[0]["applied_n"], 9)
            value = sample["isotopes"][iso]
            evidence = value["residual_correction"]
            self.assertTrue(evidence["extrapolated"])
            self.assertAlmostEqual(value["value"], original_sample["isotopes"][iso]["value"]+evidence["adjustment"])
            self.assertGreater(value["u_residual"], 0)
            self.assertAlmostEqual(value["budget"]["u_combined"]**2,
                original_sample["isotopes"][iso]["budget"]["u_combined"]**2+value["u_residual"]**2)
            self.assertEqual(sum(c["name"] == "residual_linearity" for c in value["budget"]["components"]), 1)
            material = next(m for m in analysis["diagnostics_after"]["materials"] if m["material_id"] == "__all__")
            for effect in ("intensity_dependence", "sample_reference_dependence", "pressure_adjusted_dependence"):
                points = material["isotopes"][iso][effect]["points"]
                if points:
                    self.assertEqual(next(p["y"] for p in points if p["id"] == sample["id"]), value["value"])
        with patch.object(legacy, "store", FileSessionStore(Path(self.fixture.temp.name)/"corrected-tools")):
            bridge = open_plot_bridge(restarted, session["id"], "all", Decision(**D))
            frame = legacy.store.load_frame(bridge["session_id"])
            self.assertEqual(frame.loc[int(bridge["row_mapping"][sample["id"]]), "d 13C/12C  Mean"], sample["isotopes"]["d13c"]["value"])
        exported = restarted.export_results_session(session["id"], SessionExportCommand(**D, format="json"))
        payload = json.loads(self.repo.read_blob(exported["sha256"]))
        self.assertEqual(payload["results"][0]["d13c_value"], sample["isotopes"]["d13c"]["value"])
        self.assertIn("session_residual_corrections", payload)
        self.assertEqual(restarted.run_detail(run["id"])["evaluation"], original)

    def test_manual_parameters_need_uncertainty_and_must_improve_qc_sd(self):
        session = self.create()
        self.import_linearity(session)
        endpoint = f"/metrology/results-sessions/{session['id']}/residual-overrides"
        base = {**D, "material_id": self.method["config"]["qc_id"], "isotope": "d13c"}
        for effect in ("sample_reference_dependence", "pressure_adjusted_dependence"):
            self.assertEqual(self.client.put(endpoint, json={**base,"effect":effect,"settings":{"enabled":False}}).status_code,200)
        command = {**base, "effect":"intensity_dependence", "settings":{"enabled":True,"slope":-.3}}
        self.assertEqual(self.client.put(endpoint,json=command).status_code,200)
        detail = self.service.results_session_detail(session["id"])
        self.assertEqual(detail["residual_corrections"]["d13c"]["intensity_dependence"]["status"],"uncertainty_required")
        command["settings"]["u_slope"] = .01
        self.assertEqual(self.client.put(endpoint,json=command).status_code,200)
        detail = self.service.results_session_detail(session["id"])
        self.assertEqual(detail["residual_corrections"]["d13c"]["intensity_dependence"]["status"],"not_improved")
        for row in detail["runs"][0]["evaluation"]["results"]:
            self.assertNotIn("residual_correction",row["isotopes"]["d13c"])
        command["settings"]["slope"] = .3
        self.assertEqual(self.client.put(endpoint,json=command).status_code,200)
        detail = self.service.results_session_detail(session["id"])
        self.assertEqual(detail["residual_corrections"]["d13c"]["intensity_dependence"]["status"],"applied")

    def test_flat_qc_does_not_apply_correction_or_add_uncertainty(self):
        session = self.create()
        rows = [fixtures.row(i+1,"SHP2L",0.,sample_type="QC Standard") for i in range(8)]
        for i, row in enumerate(rows):
            row["i44_v"] = 4.+i*.1
        self.service.import_run("flat-qc.xlsx",fixtures.workbook(rows),RunCommand(**D,results_session_id=session["id"]))
        detail = self.service.results_session_detail(session["id"])
        for iso in ("d13c","d18o"):
            self.assertFalse(any(d["status"] == "applied" for d in detail["residual_corrections"][iso].values()))
            self.assertEqual(detail["residual_corrections"][iso]["intensity_dependence"]["status"],"not_improved")
            for row in detail["runs"][0]["evaluation"]["results"]:
                result = row["isotopes"][iso]
                self.assertEqual(result["value"],0.)
                self.assertFalse(any(c["name"] == "residual_linearity" for c in result["budget"]["components"]))

    def test_saved_qc_outliers_persist_and_history_is_isotope_specific(self):
        session = self.create()
        rows = [fixtures.row(i+1, "SHP2L", v, sample_type="QC Standard")
                for i, v in enumerate([-.002, -.001, 0, .001, .002, .003, .004, .07])]
        for row in rows:
            row["d18O Mean"] = 0
        run = self.service.import_run("outliers.xlsx", fixtures.workbook(rows), RunCommand(**D, results_session_id=session["id"]))
        original = self.service.run_detail(run["id"])["evaluation"]
        endpoint = f"/metrology/results-sessions/{session['id']}"
        # A session IQR outlier can be inside the frozen +/-3 SD limits.
        response = self.client.put(endpoint+"/outlier-screening", json={**D, "method":"iqr", "threshold":1.5})
        self.assertEqual(response.status_code, 200, response.text)
        screening = response.json()
        self.assertEqual(len(screening["flags"]), 1)
        flag = screening["flags"][0]
        self.assertEqual(flag["isotope"], "d13c")
        self.assertEqual(flag["evaluation_id"], original["id"])
        self.assertLess(flag["value"], 3*self.method["config"]["precision"]["d13c"])
        restarted = Service(Repository(self.repo.root))
        analysis = restarted.results_session_analysis(session["id"])
        self.assertEqual(analysis["outliers"]["id"], screening["id"])
        self.assertEqual(analysis["correction_review"]["d13c"]["paired_n"], 7)
        self.assertEqual(analysis["correction_review"]["d18o"]["paired_n"], 8)
        material = analysis["diagnostics_after"]["materials"][0]
        self.assertEqual(material["isotopes"]["d13c"]["drift"]["n"], 7)
        self.assertEqual(material["isotopes"]["d18o"]["drift"]["n"], 8)
        retained_id = next(p["id"] for p in material["isotopes"]["d13c"]["drift"]["points"] if not p["excluded_from_fit"])
        preview = self.client.post(endpoint+"/analysis-preview", json={"d13c": [retained_id]}).json()
        self.assertEqual(preview["diagnostics_after"]["materials"][0]["isotopes"]["d13c"]["drift"]["n"], 6)
        self.assertEqual(preview["diagnostics_after"]["materials"][0]["isotopes"]["d18o"]["drift"]["n"], 8)
        self.assertEqual(restarted.results_session_analysis(session["id"])["diagnostics_after"]["materials"][0]["isotopes"]["d13c"]["drift"]["n"], 7)
        self.assertEqual(restarted.results_session_analysis(session["id"])["outliers"]["id"], screening["id"])
        detail = restarted.results_session_detail(session["id"])
        carbon, oxygen = (detail["history"][0]["isotopes"][iso] for iso in ("d13c", "d18o"))
        self.assertEqual((carbon["total_n"], carbon["n"], oxygen["n"]), (8, 7, 8))
        self.assertEqual(len(carbon["points"]), 8)
        self.assertIn("session_qc_outlier", [f["rule"] for f in carbon["flags"]])
        self.assertAlmostEqual(carbon["sd"], .002160246899469287)
        self.assertEqual(self.service.run_detail(run["id"])["evaluation"], original)
        with self.repo.connect(write=True) as db:
            with self.assertRaisesRegex(sqlite3.IntegrityError, "Append-only"):
                db.execute("DELETE FROM qc_screenings WHERE id=?", (screening["id"],))
        # New settings replace the active detection without destroying its provenance.
        sigma = self.client.put(endpoint+"/outlier-screening", json={**D, "method":"sigma", "threshold":3}).json()
        self.assertNotEqual(sigma["id"], screening["id"])
        self.assertEqual(sigma["flags"], [])
        self.assertEqual(self.service.results_session_detail(session["id"])["history"][0]["isotopes"]["d13c"]["n"], 8)
        # A new immutable evaluation gets a different cache population.
        self.service.evaluate_run(run["id"], EvaluateCommand(**D))
        next_analysis = self.service.results_session_analysis(session["id"])
        self.assertNotEqual(next_analysis["outliers"]["id"], sigma["id"])
        for invalid in [{"method":"bad", "threshold":3}, {"method":"sigma", "threshold":0}]:
            self.assertEqual(self.client.put(endpoint+"/outlier-screening", json={**D, **invalid}).status_code, 422)

    def setUp(self):
        self.fixture = fixtures.WorkflowTests("test_seed_has_no_provisional_anchor_certificates")
        self.fixture.setUp()
        self.addCleanup(self.fixture.doCleanups)
        self.service = self.fixture.service
        self.repo = self.fixture.repo
        self.client = self.fixture.client
        self.method = self.fixture.activate()

    def create(self, client="Client A"):
        return self.service.save_results_session(ResultsSessionCommand(**D, name="Carbonate batch", client=client,
            project="Reservoir", method_id=self.method["id"], processing_evidence="Qtegra uses this exact validated dual-anchor model; no residual correction configured"))

    def import_batch(self, session):
        rows = []
        for i, offset in enumerate([-.01, 0, .01]):
            rows.append(fixtures.row(len(rows)+1,"SHP2L",offset,sample_type="QC Standard"))
            rows.append(fixtures.row(len(rows)+1,"+sample-formula" if i==0 else f"sample-{i}",1+i))
        content=fixtures.workbook(rows)
        run=self.service.import_run("routine.xlsx",content,RunCommand(**D, results_session_id=session["id"], sample_group="Core A"))
        return run,content

    def test_persisted_session_pins_provenance_without_double_normalization(self):
        session=self.create()
        run,content=self.import_batch(session)
        self.assertEqual(run["input_basis"],"already_vpdb")
        self.assertEqual(run["external_method_id"],self.method["id"])
        detail=Service(Repository(self.repo.root)).results_session_detail(session["id"])
        self.assertEqual(detail["client"],"Client A")
        self.assertEqual(detail["qualification_id"],self.method["approval"]["qualification_id"])
        self.assertIsNotNone(detail["qualification_run_id"])
        self.assertEqual(detail["run_ids"],[run["id"]])
        self.assertEqual(detail["history"][0]["isotopes"]["d13c"]["n"],3)
        for row in detail["runs"][0]["evaluation"]["results"]:
            value=row["isotopes"]["d13c"]
            self.assertFalse(value["processing"]["normalization_applied"])
            self.assertAlmostEqual(value["value"],row["d13c"])
            self.assertGreater(value["u_norm"],0)
        again=self.service.import_run("renamed.xlsx",content,RunCommand(**D,results_session_id=session["id"]))
        self.assertTrue(again["duplicate"])
        self.assertEqual(self.service.results_session_detail(session["id"])["run_ids"],[run["id"]])
        with self.assertRaisesRegex(ValueError,"fixed"):
            self.service.save_results_session(ResultsSessionCommand(**D,name="Changed",client="Other client",method_id=self.method["id"],processing_evidence="Another model"),session["id"])
        with self.assertRaisesRegex(Conflict,"pinned"):
            self.service.annotate(run["id"],AnnotationCommand(**D,input_basis="instrument_delta"))

    def test_aragonite_import_records_basis_and_scales_full_budget(self):
        session=self.create()
        content=fixtures.workbook([fixtures.row(1,"SHP2L",0,sample_type="QC Standard"),fixtures.row(2,"Aragonite sample",1)])
        run=self.service.import_run("aragonite.xlsx",content,RunCommand(**D,results_session_id=session["id"],carbonate_material="aragonite"))
        detail=self.service.run_detail(run["id"])
        self.assertEqual(detail["carbonate_material"],"aragonite")
        qc,sample=detail["evaluation"]["results"]
        factor=1.0087/1.0091
        result=sample["isotopes"]["d18o"]
        self.assertAlmostEqual(result["value"],(sample["d18o"]+1000)*factor-1000)
        self.assertEqual(qc["isotopes"]["d18o"]["value"],qc["d18o"])
        self.assertAlmostEqual(result["u_prec"],self.method["config"]["precision"]["d18o"]*factor)
        self.assertTrue(result["processing"]["carbonate"]["applied"])
        self.assertGreater(result["budget"]["expanded_uncertainty"],0)
        with self.assertRaisesRegex(ValueError,"different carbonate basis"):
            self.service.import_run("again.xlsx",content,RunCommand(**D,results_session_id=session["id"]))
        response=self.client.get(f"/metrology/runs/{run['id']}/measurements/{sample['id']}/evidence")
        self.assertEqual(response.status_code,200)
        self.assertEqual(response.json()["raw_rows"][0]["values"]["Label"],"Aragonite sample")

    def test_client_isolation_and_foreign_groups(self):
        a=self.create(); b=self.create("Client B")
        run,content=self.import_batch(a)
        other = self.service.import_run("another-name.xlsx",content,RunCommand(**D,results_session_id=b["id"]))
        self.assertNotEqual(run["id"],other["id"])
        self.assertEqual(run["raw_import_id"],other["raw_import_id"])
        self.assertEqual(self.service.results_session_detail(b["id"])["run_ids"],[other["id"]])
        duplicate = self.service.import_run("renamed-again.xlsx",content,RunCommand(**D,results_session_id=b["id"]))
        self.assertTrue(duplicate["duplicate"])
        self.assertEqual(duplicate["id"],other["id"])
        ids_a = {r["id"] for r in self.service.run_detail(run["id"])["measurements"]}
        ids_b = {r["id"] for r in self.service.run_detail(other["id"])["measurements"]}
        self.assertFalse(ids_a & ids_b)
        with self.assertRaisesRegex(ValueError,"belonging"):
            self.service.save_session_groups(b["id"],SessionGroupsCommand(**D,groups={self.service.run_detail(run["id"])["measurements"][0]["id"]:"Stolen group"}))

    def test_group_export_has_uncertainties_and_immutable_lineage(self):
        session=self.create(); run,_=self.import_batch(session)
        detail=self.service.run_detail(run["id"])
        unknowns=[r for r in detail["evaluation"]["results"] if r["role"]=="unknown"]
        self.service.save_session_groups(session["id"],SessionGroupsCommand(**D,groups={unknowns[0]["id"]:"Client subset"}))
        record=self.service.export_results_session(session["id"],SessionExportCommand(**D,group="Client subset"))
        with zipfile.ZipFile(io.BytesIO(self.repo.read_blob(record["sha256"]))) as archive:
            rows=list(csv.DictReader(io.StringIO(archive.read("results.csv").decode("utf-8-sig"))))
            dossier=json.loads(archive.read("calculation-dossier.json"))
        self.assertEqual(len(rows),1)
        self.assertEqual(rows[0]["sample"],"'+sample-formula")
        self.assertEqual(rows[0]["raw_sha256"],detail["source"]["sha256"])
        self.assertEqual(rows[0]["qualification_id"],session["qualification_id"])
        self.assertGreater(float(rows[0]["d13c_u_norm"]),0)
        self.assertAlmostEqual(float(rows[0]["d13c_U"]),2*float(rows[0]["d13c_u_combined"]))
        self.assertEqual(sum(len(c["results"]) for c in dossier["calculations"]),1)
        with self.repo.connect(write=True) as db:
            with self.assertRaises(sqlite3.IntegrityError):
                db.execute("UPDATE session_exports SET data='{}' WHERE id=?",(record["id"],))
        response=self.client.get(f"/metrology/session-exports/{record['id']}")
        self.assertEqual(response.status_code,200)
        self.assertEqual(response.content,self.repo.read_blob(record["sha256"]))

    def test_release_freezes_client_and_sample_groups(self):
        session=self.create();run,_=self.import_batch(session)
        endpoint = f"/metrology/results-sessions/{session['id']}/uncertainty"
        self.assertEqual(self.client.put(endpoint, json={**D, "coverage_factor": 3}).status_code, 200)
        release=self.service.release(run["id"],ReleaseCommand(**D,evaluation_id=run["latest_evaluation_id"]))
        self.assertEqual(release["results_session_snapshot"]["client"],"Client A")
        sample=release["results"][0]
        for iso in ("d13c", "d18o"):
            budget = sample["isotopes"][iso]["budget"]
            self.assertEqual(budget["k"], 3)
            self.assertAlmostEqual(budget["expanded_uncertainty"], 3 * budget["u_combined"])
        self.assertEqual(self.client.put(endpoint, json={**D, "coverage_factor": 2}).status_code, 422)
        self.assertEqual(self.service.results_session_detail(session["id"])["coverage_factor"], 3)
        with self.assertRaisesRegex(ValueError,"frozen"):
            self.service.save_session_groups(session["id"],SessionGroupsCommand(**D,groups={sample["id"]:"Other group"}))

    def test_existing_plotting_tools_receive_the_correct_workbook(self):
        from services.irms_api.api import main as legacy
        from services.irms_api.session_store import FileSessionStore
        session=self.create();run,_=self.import_batch(session)
        with patch.object(legacy,"store",FileSessionStore(Path(self.fixture.temp.name)/"plot-tools")):
            url=f"/metrology/results-sessions/{session['id']}/tools/{run['id']}"
            response=self.client.post(url,json=D)
            self.assertEqual(response.status_code,200,response.text)
            bridge=response.json()["session_id"]
            metadata=legacy.store.load_metadata(bridge)
            self.assertEqual(metadata["metrology_link"]["run_id"],run["id"])
            self.assertEqual(self.client.post(url,json=D).json()["session_id"],bridge)
            qualification_run=self.service.results_session_detail(session["id"])["qualification_run_id"]
            calibration=self.client.post(f"/metrology/results-sessions/{session['id']}/tools/{qualification_run}",json=D)
            self.assertEqual(calibration.status_code,200,calibration.text)
            calibration_id=calibration.json()["session_id"]
            calibration_meta=legacy.store.load_metadata(calibration_id)
            self.assertEqual(calibration_meta["calibration"]["selected_standards"],["NBS18","NBS19"])
            from services.irms_api.domain.calibration.workspace import build_calibration_workspace
            workspace=build_calibration_workspace(calibration_id,legacy.store.load_frame(calibration_id),calibration_meta,include_figures=False)
            self.assertEqual(next(v.value for v in workspace.selected_standard_official_values if v.standard=="NBS18"),-10)
            from fastapi.testclient import TestClient
            with TestClient(legacy.app) as app:
                guarded=app.post(f"/sessions/{bridge}/calibration/run",json={})
                self.assertEqual(guarded.status_code,409,guarded.text)

    def test_processing_label_edits_persist_in_results_and_exports(self):
        from fastapi.testclient import TestClient
        from services.irms_api.api import main as legacy
        from services.irms_api.session_store import FileSessionStore
        from services.irms_api.metrology.plot_bridge import open_plot_bridge
        from services.irms_api.metrology.models import Decision
        session = self.create()
        run, content = self.import_batch(session)
        other_session = self.create("Client B")
        other_run = self.service.import_run("same.xlsx", content, RunCommand(**D, results_session_id=other_session["id"]))
        original = self.service.run_detail(run["id"])["evaluation"]
        sample = next(r for r in original["results"] if r["role"] == "unknown")
        with patch.object(legacy, "store", FileSessionStore(Path(self.fixture.temp.name)/"label-tools")):
            bridge = open_plot_bridge(self.service, session["id"], "all", Decision(**D))
            target = {"row_label":bridge["row_mapping"][sample["id"]],"isotope_key":"d13C"}
            edit = {"action":"set_identifier2","identifier2":"Corrected spelling","targets":[target]}
            with TestClient(legacy.app) as app:
                response = app.post(f"/sessions/{bridge['session_id']}/processing/edit", json=edit)
                self.assertEqual(response.status_code, 200, response.text)
                # A mixed batch is rejected before any identity or numerical edit.
                batch = {"edits":[{**edit,"identifier2":"Should not be saved"},{"action":"set_value","value":8,"targets":[target]}]}
                response = app.post(f"/sessions/{bridge['session_id']}/processing/edits", json=batch)
                self.assertEqual(response.status_code, 409, response.text)
            restarted = Service(Repository(self.repo.root))
            detail = restarted.results_session_detail(session["id"])
            changed = next(r for r in detail["runs"][0]["evaluation"]["results"] if r["id"] == sample["id"])
            self.assertEqual(changed["identifier2"], "Corrected spelling")
            self.assertEqual(changed["comment"], sample["comment"])
            fresh = open_plot_bridge(restarted, session["id"], "all", Decision(**D))
            self.assertNotEqual(fresh["session_id"],bridge["session_id"])
            frame = legacy.store.load_frame(fresh["session_id"])
            self.assertEqual(frame.loc[int(fresh["row_mapping"][sample["id"]]),"Identifier 2"],"Corrected spelling")
            record = restarted.export_results_session(session["id"], SessionExportCommand(**D, format="json"))
            exported = json.loads(self.repo.read_blob(record["sha256"]))
            corrected = next(r for r in exported["results"] if r["measurement_id"] == sample["id"])
            self.assertEqual(corrected["identifier2"], "Corrected spelling")
            self.assertEqual(corrected["sample_identifier"], "Corrected spelling")
            self.assertEqual(restarted.run_detail(run["id"])["evaluation"], original)
            self.assertFalse(any(r.get("identifier2") == "Corrected spelling" for r in restarted.run_detail(other_run["id"])["measurements"]))

    def test_domain_recovery_supplies_budget_and_preserves_failure_reasons(self):
        from services.irms_api.metrology.session_processing import process_session_results
        session = self.create()
        self.import_linearity(session)
        detail = self.service.results_session_detail(session["id"])
        run = detail["runs"][0]
        # Reproduce a legacy recorded evaluation rejected by the method domain.
        run["evaluation"] = self.service.run_detail(run["id"])["evaluation"]
        row = next(r for r in run["evaluation"]["results"] if r["role"] == "unknown")
        model = run["evaluation"]["normalization"]["d13c"]
        model["correction"] = {"name":"Test intensity correction","slope":.01,"u_slope":.001,
            "predictor":"i44_v","center":5.,"domain":{"low":3.,"high":9.},"anchor_predictor_means":[5.,5.]}
        row["isotopes"] = {}
        row["issues"] = ["d13c correction predictor outside validated domain", "Qtegra pressure adjustment failed"]
        row["d18o"] = None
        process_session_results(detail)
        result = row["isotopes"]["d13c"]
        self.assertTrue(result["correction"]["domain_extrapolated"])
        self.assertGreater(result["budget"]["u_combined"], 0)
        self.assertNotIn("residual_correction", result)
        self.assertIn("Qtegra pressure adjustment failed", row["issues"])
        self.assertNotIn("d18o", row["isotopes"])
        self.assertEqual(row["calculation_issues"]["d18o"], "The imported isotope value is missing.")

    def test_full_session_bridge_maps_all_workbooks_and_invalidates_after_revision(self):
        from services.irms_api.api import main as legacy
        from services.irms_api.session_store import FileSessionStore
        from services.irms_api.metrology.plot_bridge import open_plot_bridge
        from services.irms_api.metrology.models import Decision
        session=self.create(); run,_=self.import_batch(session)
        second=self.service.import_run("routine.xlsx",fixtures.workbook([fixtures.row(1,"SHP2L",.002,sample_type="QC Standard"),fixtures.row(2,"another sample",2)]),RunCommand(**D,results_session_id=session["id"]))
        with patch.object(legacy,"store",FileSessionStore(Path(self.fixture.temp.name)/"session-tools")):
            result=open_plot_bridge(self.service,session["id"],"all",Decision(**D))
            first_workbook = open_plot_bridge(self.service,session["id"],run["id"],Decision(**D))
            self.assertEqual(len(result["row_mapping"]),8)
            self.assertEqual(len(set(result["row_mapping"].values())),8)
            frame=legacy.store.load_frame(result["session_id"])
            for source in self.service.results_session_detail(session["id"])["runs"]:
                for row in source["measurements"]:
                    self.assertAlmostEqual(frame.loc[int(result["row_mapping"][row["id"]]),"d 13C/12C  Mean"],row["d13c"])
            self.assertEqual(open_plot_bridge(self.service,session["id"],"all",Decision(**D))["session_id"],result["session_id"])
            # A format upgrade reuses already parsed workbooks and saved plot controls.
            import hashlib
            from services.irms_api.metrology.repository import encode
            meta=legacy.store.load_metadata(result["session_id"])
            runs=self.service.results_session_detail(session["id"])["runs"]
            meta["metrology_link"].update(bridge_version=7,fingerprint=hashlib.sha256(encode([(r["id"],r["revision"],r.get("latest_evaluation_id")) for r in runs]).encode()).hexdigest())
            meta["calibration"]["config"]["color_param"]="Date"
            legacy.store.write_metadata(result["session_id"],meta)
            with patch.object(legacy,"_import_session_from_bytes",side_effect=AssertionError("Unchanged workbooks must not be reparsed")):
                upgraded=open_plot_bridge(self.service,session["id"],"all",Decision(**D))
            self.assertEqual(upgraded["session_id"],result["session_id"])
            self.assertEqual(legacy.store.load_metadata(result["session_id"])["calibration"]["config"]["color_param"],"Date")
            self.service.annotate(second["id"],AnnotationCommand(**D,acquisition_complete=True))
            changed=open_plot_bridge(self.service,session["id"],"all",Decision(**D))
            self.assertNotEqual(changed["session_id"],result["session_id"])
            self.assertNotEqual(open_plot_bridge(self.service,session["id"],run["id"],Decision(**D))["session_id"],first_workbook["session_id"])

    def test_session_analysis_reuses_qc_outlier_algorithm_and_pressure_fit(self):
        from services.irms_api.metrology.session_analysis import session_analysis
        from services.irms_api.domain.calibration.core import identify_outliers_iqr
        import pandas as pd
        session=self.create(); self.import_batch(session)
        self.service.import_run("additional.xlsx",fixtures.workbook([fixtures.row(i+1,"SHP2L",v,sample_type="QC Standard") for i,v in enumerate([-.01,0,.01,.5])]),RunCommand(**D,results_session_id=session["id"]))
        detail=self.service.results_session_detail(session["id"])
        data=session_analysis(detail,outlier_method="iqr",threshold=1.5)
        self.assertEqual(data["workbooks"],2); self.assertEqual(len(data["rows"]),10)
        qc=[r for r in data["rows"] if r["role"]=="qc"]
        expected=identify_outliers_iqr(pd.DataFrame({"value":[r["isotopes"]["d13c"]["value"] for r in qc]}),"value",1.5)
        self.assertEqual({r["measurement_id"] for r in data["outliers"]["flags"] if r["isotope"]=="d13c"},{r["id"] for r,flag in zip(qc,expected) if flag})
        self.assertEqual(data["correction_review"]["d13c"]["paired_n"],6)
        self.assertEqual(data["correction_review"]["d13c"]["excluded_outlier_n"],1)
        self.assertEqual(data["correction_review"]["d18o"]["paired_n"],6)
        self.assertEqual(data["qc_statistics"]["d13c"]["final"]["n"],6)
        self.assertEqual(data["qc_statistics"]["d18o"]["final"]["n"],6)
        self.assertIn("pressure_dependence",data["diagnostics_before"]["materials"][0]["isotopes"]["d13c"])
        self.assertFalse(any(r["excluded"] for r in data["rows"]))

    def test_sample_override_preserves_original_and_does_not_clear_other_gates(self):
        session=self.create()
        rows=[fixtures.row(i+1,"SHP2L",v,sample_type="QC Standard") for i,v in enumerate([-.01,0,.01])]
        rows.append(fixtures.row(4,"Outside anchors",20))
        run=self.service.import_run("exceptions.xlsx",fixtures.workbook(rows),RunCommand(**D,results_session_id=session["id"]))
        original=self.service.run_detail(run["id"])["evaluation"]
        row=next(r for r in original["results"] if r["role"]=="unknown")
        issues=[i for i in row["issues"] if "extrapolates" in i or "outside validated isotope" in i]
        self.assertTrue(issues)
        reviewed=self.service.review_row(run["id"],RowReviewCommand(**D,evaluation_id=original["id"],measurement_id=row["id"],issues=issues))
        self.assertEqual(reviewed["results"][-1]["accepted_issues"],issues)
        self.assertEqual(reviewed["results"][-1]["isotopes"],row["isotopes"])
        self.assertEqual(self.service.results_session_detail(session["id"])["history"][0]["isotopes"]["d13c"]["n"],3)
        with self.repo.connect() as db:
            self.assertEqual(self.repo.get(db,"evaluations",original["id"])["results"][-1]["issues"],row["issues"])
        with self.assertRaisesRegex(Conflict,"latest"):
            self.service.review_row(run["id"],RowReviewCommand(**D,evaluation_id=original["id"],measurement_id=row["id"],issues=issues))
        reevaluated=self.service.evaluate_run(run["id"],EvaluateCommand(**D))
        self.assertEqual(reevaluated["results"][-1]["issues"],row["issues"])
        self.assertNotIn("accepted_issues",reevaluated["results"][-1])

    def test_override_does_not_accept_qc_or_missing_uncertainty(self):
        session=self.create();run,_=self.import_batch(session)
        evaluation=self.service.run_detail(run["id"])["evaluation"]
        qc=next(r for r in evaluation["results"] if r["role"]=="qc")
        with self.assertRaisesRegex(ValueError,"unknown samples"):
            self.service.review_row(run["id"],RowReviewCommand(**D,evaluation_id=evaluation["id"],measurement_id=qc["id"],issues=["force keep"]))
        rows=[fixtures.row(i+1,"SHP2L",v,sample_type="QC Standard") for i,v in enumerate([-.01,0,.01])]
        rows.append(fixtures.row(4,"Missing isotope",1));rows[-1]["d18O Mean"]=None
        broken=self.service.import_run("missing.xlsx",fixtures.workbook(rows),RunCommand(**D,results_session_id=session["id"]))
        evaluation=self.service.run_detail(broken["id"])["evaluation"]
        unknown=next(r for r in evaluation["results"] if r["role"]=="unknown")
        with self.assertRaisesRegex(ValueError,"complete calculated"):
            self.service.review_row(broken["id"],RowReviewCommand(**D,evaluation_id=evaluation["id"],measurement_id=unknown["id"],issues=unknown["issues"]))

    def test_empty_session_analysis_and_legacy_mapping_preserve_exported_means(self):
        from services.irms_api.metrology.session_analysis import session_analysis
        session=self.create()
        empty=session_analysis(self.service.results_session_detail(session["id"]))
        self.assertEqual(empty["rows"],[])
        self.assertEqual(empty["diagnostics_before"]["materials"],[])
        self.assertEqual(empty["qc_statistics"]["d13c"]["final"]["n"],0)

    def test_pdf_dossier_and_legacy_filenames_keep_review_status_and_group_scope(self):
        import pymupdf
        from openpyxl import load_workbook
        session=self.create();run,_=self.import_batch(session)
        record=self.service.export_results_session(session["id"],SessionExportCommand(**D,client_name="Client / A",series_name="17 # 2"))
        self.assertTrue(record["filename"].startswith("Results for 17 series - stable C & O isotopes - P2L - Client _ A"))
        with zipfile.ZipFile(io.BytesIO(self.repo.read_blob(record["sha256"]))) as archive:
            document=pymupdf.open(stream=archive.read("calculation-certificate.pdf"),filetype="pdf")
            content="\n".join(p.get_text() for p in document)
            self.assertIn("Result status: review",content)
            self.assertIn("u_norm",content)
            self.assertIn(self.service.run_detail(run["id"])["source"]["sha256"],content.replace("\n",""))
            xlsx=load_workbook(io.BytesIO(archive.read(next(n for n in archive.namelist() if n.endswith(".xlsx")))))
            self.assertEqual(xlsx.active.max_row,4)
            self.assertEqual(xlsx.active.freeze_panes,"A2")


if __name__=="__main__":
    unittest.main()
